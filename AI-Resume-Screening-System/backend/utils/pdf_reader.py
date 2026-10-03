"""
Document File Parser (PDF and DOCX extraction)
Supports pdfplumber, PyPDF2, and python-docx.
"""
import os
import io
import logging

logger = logging.getLogger(__name__)

def extract_text_from_pdf(file_path_or_stream) -> str:
    """Extracts raw text from a PDF file using pdfplumber with PyPDF2 fallback."""
    text = ""
    
    # Try pdfplumber first for high precision
    try:
        import pdfplumber
        with pdfplumber.open(file_path_or_stream) as pdf:
            pages_text = []
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    pages_text.append(t)
            text = "\n".join(pages_text)
            if text.strip():
                return text
    except Exception as e:
        logger.warning(f"pdfplumber extraction failed: {e}. Trying PyPDF2...")

    # Fallback to PyPDF2
    try:
        import PyPDF2
        if isinstance(file_path_or_stream, str):
            with open(file_path_or_stream, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                pages_text = [p.extract_text() or "" for p in reader.pages]
                text = "\n".join(pages_text)
        else:
            reader = PyPDF2.PdfReader(file_path_or_stream)
            pages_text = [p.extract_text() or "" for p in reader.pages]
            text = "\n".join(pages_text)
        if text.strip():
            return text
    except Exception as e:
        logger.error(f"PyPDF2 extraction failed: {e}")

    return text

def extract_text_from_docx(file_path_or_stream) -> str:
    """Extracts raw text from a DOCX file using python-docx."""
    try:
        import docx
        doc = docx.Document(file_path_or_stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except Exception as e:
        logger.error(f"python-docx extraction failed: {e}")
        return ""

def extract_document_text(file_path: str) -> str:
    """Detects extension and routes to appropriate extractor."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported document format: {ext}")
